"""
modules/agreement_template.py
표준 업무협약서 (2자용/3자용) DOCX 생성 — HWPX 원본 충실 재현

폰트  : 휴먼명조(본문) / HY울릉도M(헤딩) / 휴먼고딕(서명기관명)
테두리: 외곽 0.7mm / 내부 0.12mm (셀별 방향 개별 적용)
배경색: 없음
서명란: 본문 셀 내 중첩 표 (테두리 없음, 우측 정렬 103.24mm)
헤딩  : 3열 표 (navy #1C3D62 배경 / empty / 아래선)
"""

from __future__ import annotations
from datetime import date
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).parent))

from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from doc_builder import OUTPUTS_DIR, save_docx

# ══════════════════════════════════════════════
# 상수
# ══════════════════════════════════════════════

FONT_BODY = "휴먼명조"   # 표 본문
FONT_HEAD = "HY울릉도M"  # 제목 헤딩
FONT_SIGN = "휴먼고딕"   # 서명 기관명
FALLBACK  = "맑은 고딕"

PT_TITLE  = 15   # 협약서명 제목
PT_LABEL  = 13   # 협약당사자/협약내용/기관명 라벨
PT_BODY   = 13   # 지원내용·협약기간·본문 텍스트
PT_SMALL  = 10   # 협약번호
PT_SIGN   = 14   # 서명 기관명
PT_HEAD_L = 15   # 헤딩 "공통 N"
PT_HEAD_R = 16   # 헤딩 "표준협약서(안)"

L_MARGIN = Mm(21.17)
R_MARGIN = Mm(19.49)
T_MARGIN = Mm(19.99)
B_MARGIN = Mm(14.99)

# 본문 표 4열 너비 (HWPUNIT 실측)
COL_A = Mm(9.74)    # 라벨 열
COL_B = Mm(23.71)   # 세부 라벨
COL_C = Mm(83.61)   # 내용 열
COL_D = Mm(49.68)   # 협약번호/우측

# 헤딩 표 3열 너비
HEAD_A = Mm(17.55)
HEAD_B = Mm(1.99)
HEAD_C = Mm(101.94)

# 서명 표 2열 너비
SIGN_ORG   = Mm(88.85)
SIGN_STAMP = Mm(14.39)

# 테두리 굵기 (OOXML sz = 1/8 pt)
SZ_THICK = 16    # 2pt ≈ 0.7mm  (외곽)
SZ_THIN  = 3     # 0.375pt ≈ 0.12mm  (내부)
SZ_MED   = 14    # 1.75pt ≈ 0.6mm  (헤딩 하단선)


# ══════════════════════════════════════════════
# XML 헬퍼
# ══════════════════════════════════════════════

def _emu2twips(emu) -> str:
    return str(int(emu / 914400 * 1440))


def _font(run, name: str = FONT_BODY, size: float = PT_BODY,
          bold: bool = False, color: str | None = None) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(
            int(color[0:2], 16), int(color[2:4], 16), int(color[4:6], 16)
        )
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(attr), name)


def _shade(cell, hex_fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _valign(cell, val: str = "center") -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    old = tcPr.find(qn("w:vAlign"))
    if old is not None:
        tcPr.remove(old)
    vAlign = OxmlElement("w:vAlign")
    vAlign.set(qn("w:val"), val)
    tcPr.append(vAlign)


def _no_spacing(para) -> None:
    pf = para.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after  = Pt(0)


def _no_auto_space(para) -> None:
    pPr = para._p.get_or_add_pPr()
    for tag in ("w:autoSpaceDE", "w:autoSpaceDN"):
        old = pPr.find(qn(tag))
        if old is not None:
            pPr.remove(old)
        el = OxmlElement(tag)
        el.set(qn("w:val"), "0")
        pPr.append(el)


def _set_row_height(row, height_mm: float, exact: bool = False) -> None:
    trPr = row._tr.get_or_add_trPr()
    old = trPr.find(qn("w:trHeight"))
    if old is not None:
        trPr.remove(old)
    trH = OxmlElement("w:trHeight")
    trH.set(qn("w:val"), _emu2twips(Mm(height_mm)))
    trH.set(qn("w:hRule"), "exact" if exact else "atLeast")
    trPr.append(trH)


def _set_col_widths(tbl, widths: list) -> None:
    tblEl = tbl._tbl
    old = tblEl.find(qn("w:tblGrid"))
    if old is not None:
        tblEl.remove(old)
    grid = OxmlElement("w:tblGrid")
    for w in widths:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), _emu2twips(w))
        grid.append(gc)
    tblPr = tblEl.find(qn("w:tblPr"))
    idx = list(tblEl).index(tblPr) + 1 if tblPr is not None else 0
    tblEl.insert(idx, grid)
    for row in tbl.rows:
        for ci, cell in enumerate(row.cells):
            if ci < len(widths):
                tcPr = cell._tc.get_or_add_tcPr()
                old_w = tcPr.find(qn("w:tcW"))
                if old_w is not None:
                    tcPr.remove(old_w)
                tcW = OxmlElement("w:tcW")
                tcW.set(qn("w:w"), _emu2twips(widths[ci]))
                tcW.set(qn("w:type"), "dxa")
                tcPr.append(tcW)


def _set_cell_borders(cell,
                       top: int | None = None,
                       left: int | None = None,
                       bottom: int | None = None,
                       right: int | None = None,
                       color: str = "000000") -> None:
    """셀 방향별 테두리. sz: OOXML 1/8pt 단위. None = NONE."""
    tcPr = cell._tc.get_or_add_tcPr()
    old = tcPr.find(qn("w:tcBorders"))
    if old is not None:
        tcPr.remove(old)
    borders = OxmlElement("w:tcBorders")
    for side, sz in [("top", top), ("left", left), ("bottom", bottom), ("right", right)]:
        el = OxmlElement(f"w:{side}")
        if sz is None:
            el.set(qn("w:val"), "none")
        else:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(sz))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), color)
        borders.append(el)
    tcPr.append(borders)


def _disable_table_borders(tbl) -> None:
    tblPr = tbl._tbl.find(qn("w:tblPr"))
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl._tbl.insert(0, tblPr)
    old = tblPr.find(qn("w:tblBorders"))
    if old is not None:
        tblPr.remove(old)
    bd = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "none")
        bd.append(el)
    tblPr.append(bd)


def _set_table_cell_margins(tbl, left_mm: float = 1.80, right_mm: float = 1.80,
                              top_mm: float = 0.50, bottom_mm: float = 0.50) -> None:
    tblPr = tbl._tbl.find(qn("w:tblPr"))
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl._tbl.insert(0, tblPr)
    old = tblPr.find(qn("w:tblCellMar"))
    if old is not None:
        tblPr.remove(old)
    tcMar = OxmlElement("w:tblCellMar")
    for side, mm in [("top", top_mm), ("left", left_mm),
                      ("bottom", bottom_mm), ("right", right_mm)]:
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(int(Mm(mm) / 914400 * 1440)))
        el.set(qn("w:type"), "dxa")
        tcMar.append(el)
    tblPr.append(tcMar)


def _set_table_width(tbl, width_emu) -> None:
    tblPr = tbl._tbl.find(qn("w:tblPr"))
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl._tbl.insert(0, tblPr)
    old = tblPr.find(qn("w:tblW"))
    if old is not None:
        tblPr.remove(old)
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:w"), _emu2twips(width_emu))
    tblW.set(qn("w:type"), "dxa")
    tblPr.append(tblW)


def _write(cell, text: str, font: str = FONT_BODY, size: float = PT_BODY,
           bold: bool = False, align=WD_ALIGN_PARAGRAPH.LEFT,
           color: str | None = None, valign: str = "center") -> None:
    cell.text = ""
    para = cell.paragraphs[0]
    para.alignment = align
    _no_spacing(para)
    _no_auto_space(para)
    if text:
        run = para.add_run(text)
        _font(run, name=font, size=size, bold=bold, color=color)
    _valign(cell, valign)


def _write_multiline(cell, lines: list[str], font: str = FONT_BODY,
                     size: float = PT_BODY, align=WD_ALIGN_PARAGRAPH.LEFT,
                     valign: str = "center") -> None:
    cell.text = ""
    for i, line in enumerate(lines):
        para = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        para.alignment = align
        _no_spacing(para)
        _no_auto_space(para)
        run = para.add_run(line)
        _font(run, name=font, size=size)
    _valign(cell, valign)


def _add_highlight(run, color: str = "yellow") -> None:
    """run에 형광펜 색상 적용."""
    rPr = run._r.get_or_add_rPr()
    old = rPr.find(qn("w:highlight"))
    if old is not None:
        rPr.remove(old)
    hl = OxmlElement("w:highlight")
    hl.set(qn("w:val"), color)
    rPr.append(hl)


def _write_org_content(cell, rep: str, reg: str, addr: str,
                       font: str = FONT_BODY, size: float = PT_BODY) -> None:
    """
    기관 대표자/사업자번호/주소를 작성.
    라벨 부분("‣ 대표자: " 등)에 노란 형광펜 적용.
    """
    items = [
        ("‣ 대표자: ",        rep),
        ("‣ 사업자등록번호: ", reg),
        ("‣ 주  소: ",        addr),
    ]
    cell.text = ""
    _valign(cell, "center")
    for i, (label, value) in enumerate(items):
        para = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        _no_spacing(para)
        _no_auto_space(para)
        r_lbl = para.add_run(label)
        _font(r_lbl, name=font, size=size)
        _add_highlight(r_lbl, "yellow")
        if value:
            r_val = para.add_run(value)
            _font(r_val, name=font, size=size)


# ══════════════════════════════════════════════
# 중첩 서명 표 XML 빌더
# ══════════════════════════════════════════════

def _make_sign_cell_xml(text: str, width_emu, is_org: bool):
    tc = OxmlElement("w:tc")

    tcPr = OxmlElement("w:tcPr")
    tcW = OxmlElement("w:tcW")
    tcW.set(qn("w:w"), _emu2twips(width_emu))
    tcW.set(qn("w:type"), "dxa")
    tcPr.append(tcW)
    tcBd = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "none")
        tcBd.append(el)
    tcPr.append(tcBd)
    vAlign = OxmlElement("w:vAlign")
    vAlign.set(qn("w:val"), "center")
    tcPr.append(vAlign)
    tc.append(tcPr)

    p = OxmlElement("w:p")
    pPr = OxmlElement("w:pPr")
    jc = OxmlElement("w:jc")
    jc.set(qn("w:val"), "center")
    pPr.append(jc)
    sp = OxmlElement("w:spacing")
    sp.set(qn("w:before"), "0")
    sp.set(qn("w:after"), "0")
    pPr.append(sp)
    p.append(pPr)

    if text:
        r = OxmlElement("w:r")
        rPr = OxmlElement("w:rPr")
        rFonts = OxmlElement("w:rFonts")
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rFonts.set(qn(attr), FONT_SIGN)
        rPr.append(rFonts)
        if is_org:
            rPr.append(OxmlElement("w:b"))
        sz_el = OxmlElement("w:sz")
        sz_el.set(qn("w:val"), str(int(PT_SIGN * 2)))
        rPr.append(sz_el)
        szCs = OxmlElement("w:szCs")
        szCs.set(qn("w:val"), str(int(PT_SIGN * 2)))
        rPr.append(szCs)
        r.append(rPr)
        t = OxmlElement("w:t")
        t.text = text
        if text.startswith(" ") or text.endswith(" "):
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        r.append(t)
        p.append(r)

    tc.append(p)
    return tc


def _build_nested_sign_xml(orgs: list[str]):
    """서명란 중첩 표 XML (w:tbl). 테두리 없음, 우측 정렬, 103.24mm."""
    total_emu = SIGN_ORG + SIGN_STAMP

    tbl = OxmlElement("w:tbl")

    tblPr = OxmlElement("w:tblPr")
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:w"), _emu2twips(total_emu))
    tblW.set(qn("w:type"), "dxa")
    tblPr.append(tblW)
    jc = OxmlElement("w:jc")
    jc.set(qn("w:val"), "right")
    tblPr.append(jc)
    tblBd = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "none")
        tblBd.append(el)
    tblPr.append(tblBd)
    tbl.append(tblPr)

    grid = OxmlElement("w:tblGrid")
    gc1 = OxmlElement("w:gridCol")
    gc1.set(qn("w:w"), _emu2twips(SIGN_ORG))
    gc2 = OxmlElement("w:gridCol")
    gc2.set(qn("w:w"), _emu2twips(SIGN_STAMP))
    grid.append(gc1)
    grid.append(gc2)
    tbl.append(grid)

    for org_name in orgs:
        tr = OxmlElement("w:tr")
        trPr = OxmlElement("w:trPr")
        trH = OxmlElement("w:trHeight")
        trH.set(qn("w:val"), _emu2twips(Mm(9.93)))
        trH.set(qn("w:hRule"), "atLeast")
        trPr.append(trH)
        tr.append(trPr)
        tr.append(_make_sign_cell_xml(org_name, SIGN_ORG, is_org=True))
        tr.append(_make_sign_cell_xml("印", SIGN_STAMP, is_org=False))
        tbl.append(tr)

    return tbl


# ══════════════════════════════════════════════
# 헤딩 표 빌더
# ══════════════════════════════════════════════

def _build_heading_table(doc: Document, heading_num: int = 3) -> None:
    """
    '공통 N  표준협약서(안)' 제목 헤딩을 3열 표로 생성.
    열0: "공통 N" navy bg, HY울릉도M 15pt 흰색
    열1: empty spacer
    열2: " 표준협약서(안)" 아래선만 0.6mm
    """
    tbl = doc.add_table(rows=1, cols=3)
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    _disable_table_borders(tbl)
    _set_col_widths(tbl, [HEAD_A, HEAD_B, HEAD_C])
    _set_table_width(tbl, HEAD_A + HEAD_B + HEAD_C)
    _set_table_cell_margins(tbl, 0, 0, 0, 0)
    _set_row_height(tbl.rows[0], 9.34)

    c0 = tbl.cell(0, 0)
    _write(c0, f"공통 {heading_num}",
           font=FONT_HEAD, size=PT_HEAD_L, bold=False,
           align=WD_ALIGN_PARAGRAPH.CENTER, color="FFFFFF")
    _shade(c0, "1C3D62")
    _set_cell_borders(c0)

    c1 = tbl.cell(0, 1)
    c1.text = ""
    _set_cell_borders(c1)

    c2 = tbl.cell(0, 2)
    _write(c2, " 표준협약서(안)",
           font=FONT_HEAD, size=PT_HEAD_R, bold=False,
           align=WD_ALIGN_PARAGRAPH.LEFT)
    _set_cell_borders(c2, bottom=SZ_MED)


# ══════════════════════════════════════════════
# 표지 표 빌더
# ══════════════════════════════════════════════

def _build_cover_table(doc: Document, data: dict, party_type: str = '2way') -> None:
    """
    2자용: 7행×4열 / 3자용: 8행×4열
    행0~1: 제목(cs=3 rs=2) + 협약번호
    행2~N: 협약당사자 (당사자 수만큼)
    행N+1~N+2: 협약내용 (지원내용/협약기간)
    행 마지막: 본문+서명 (cs=4, 중첩 서명 표 포함)
    """
    is_2way    = party_type == '2way'
    party_rows = 2 if is_2way else 3
    ROWS = 2 + party_rows + 2 + 1  # 제목2 + 당사자N + 내용2 + 본문1

    pn  = data.get("project_name",    "000 지원사업")
    cw  = data.get("center_wide",     "00광역자활센터")
    cl  = data.get("center_local",    "00지역자활센터")
    ce  = data.get("enterprise_name", "00자활기업")
    rw  = data.get("rep_wide",        "")
    rl  = data.get("rep_local",       "")
    re_ = data.get("rep_enterprise",  "")
    rgw = data.get("reg_wide",        "")
    rgl = data.get("reg_local",       "")
    rge = data.get("reg_enterprise",  "")
    aw  = data.get("addr_wide",       "")
    al  = data.get("addr_local",      "")
    ae  = data.get("addr_enterprise", "")
    sc  = data.get("support_content", "인건비, 자산취득, 시설개선 지원" if not is_2way else "사업비 지원")
    ps  = data.get("period_start",    "2025.00.00.")
    pe  = data.get("period_end",      "2025.00.00.")
    cno = data.get("contract_no",     "")
    sd  = data.get("sign_date",       "2025.  00.  00.")

    tbl = doc.add_table(rows=ROWS, cols=4)
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    _disable_table_borders(tbl)
    _set_col_widths(tbl, [COL_A, COL_B, COL_C, COL_D])
    _set_table_width(tbl, COL_A + COL_B + COL_C + COL_D)
    _set_table_cell_margins(tbl)
    # 열 너비 고정 (Word 자동 재분배 방지)
    tblPr_main = tbl._tbl.find(qn("w:tblPr"))
    tblLayout = OxmlElement("w:tblLayout")
    tblLayout.set(qn("w:type"), "fixed")
    tblPr_main.append(tblLayout)

    # ── 행0~1: 제목 + 협약번호 ─────────────────────────────────
    title_cell = tbl.cell(0, 0).merge(tbl.cell(1, 2))
    _write(title_cell,
           f"{pn} 협약서({'2' if is_2way else '3'}자용)",
           font=FONT_BODY, size=PT_TITLE, bold=True,
           align=WD_ALIGN_PARAGRAPH.CENTER)
    _no_auto_space(title_cell.paragraphs[0])
    # 제목 셀: 외곽(L,T=THICK) + 내부 구분선 없음(R,B=NONE)
    _set_cell_borders(title_cell,
                      top=SZ_THICK, left=SZ_THICK,
                      bottom=None, right=None)
    _set_row_height(tbl.rows[0], 13.04)

    no_lbl = tbl.cell(0, 3)
    _write(no_lbl, "협약번호 :",
           font=FONT_BODY, size=PT_SMALL, bold=True,
           align=WD_ALIGN_PARAGRAPH.CENTER)
    # 협약번호 라벨: 외곽(T,R=THICK) + 내부 구분선 없음(L,B=NONE)
    _set_cell_borders(no_lbl,
                      top=SZ_THICK, right=SZ_THICK,
                      left=None, bottom=None)

    no_val = tbl.cell(1, 3)
    _write(no_val, cno, font=FONT_BODY, size=PT_SMALL,
           align=WD_ALIGN_PARAGRAPH.LEFT)
    # 협약번호 값: 외곽(R=THICK, B=THIN) + 내부 구분선 없음(L,T=NONE)
    _set_cell_borders(no_val,
                      right=SZ_THICK, bottom=SZ_THIN,
                      top=None, left=None)
    _set_row_height(tbl.rows[1], 6.52)

    # ── 협약당사자 행 ──────────────────────────────────────────
    r_start = 2
    r_end   = r_start + party_rows - 1

    party_lbl = tbl.cell(r_start, 0).merge(tbl.cell(r_end, 0))
    _write(party_lbl, "협약당사자",
           font=FONT_BODY, size=PT_LABEL, bold=True,
           align=WD_ALIGN_PARAGRAPH.CENTER)
    _set_cell_borders(party_lbl,
                      left=SZ_THICK,
                      top=SZ_THIN, right=SZ_THIN, bottom=SZ_THIN)

    orgs_info = [
        (cw, rw, rgw, aw),
        (cl, rl, rgl, al),
    ]
    if not is_2way:
        orgs_info.append((ce, re_, rge, ae))

    for i, (org_name, rep, reg, addr) in enumerate(orgs_info):
        ri = r_start + i
        org_cell = tbl.cell(ri, 1)
        _write(org_cell, org_name,
               font=FONT_BODY, size=PT_LABEL, bold=True,
               align=WD_ALIGN_PARAGRAPH.CENTER)
        _set_cell_borders(org_cell,
                          top=SZ_THIN, left=SZ_THIN,
                          bottom=SZ_THIN, right=SZ_THIN)

        content_cell = tbl.cell(ri, 2).merge(tbl.cell(ri, 3))
        _write_org_content(content_cell, rep, reg, addr)
        _set_cell_borders(content_cell,
                          right=SZ_THICK,
                          top=SZ_THIN, left=SZ_THIN, bottom=SZ_THIN)
        _set_row_height(tbl.rows[ri], 37.47)

    # ── 협약내용 행 ────────────────────────────────────────────
    r_content = r_end + 1

    content_lbl = tbl.cell(r_content, 0).merge(tbl.cell(r_content + 1, 0))
    _write(content_lbl, "협약내용",
           font=FONT_BODY, size=PT_LABEL, bold=True,
           align=WD_ALIGN_PARAGRAPH.CENTER)
    _set_cell_borders(content_lbl,
                      left=SZ_THICK,
                      top=SZ_THIN, right=SZ_THIN, bottom=SZ_THIN)

    si_lbl = tbl.cell(r_content, 1)
    _write(si_lbl, "지원내용",
           font=FONT_BODY, size=PT_LABEL, bold=True,
           align=WD_ALIGN_PARAGRAPH.CENTER)
    _set_cell_borders(si_lbl,
                      top=SZ_THIN, left=SZ_THIN, bottom=SZ_THIN, right=SZ_THIN)

    si_val = tbl.cell(r_content, 2).merge(tbl.cell(r_content, 3))
    _write(si_val, sc, font=FONT_BODY, size=PT_BODY,
           align=WD_ALIGN_PARAGRAPH.LEFT)
    _set_cell_borders(si_val,
                      right=SZ_THICK,
                      top=SZ_THIN, left=SZ_THIN, bottom=SZ_THIN)
    _set_row_height(tbl.rows[r_content], 14.66)

    pd_lbl = tbl.cell(r_content + 1, 1)
    _write(pd_lbl, "협약기간",
           font=FONT_BODY, size=PT_LABEL, bold=True,
           align=WD_ALIGN_PARAGRAPH.CENTER)
    _set_cell_borders(pd_lbl,
                      top=SZ_THIN, left=SZ_THIN, bottom=SZ_THIN, right=SZ_THIN)

    pd_val = tbl.cell(r_content + 1, 2).merge(tbl.cell(r_content + 1, 3))
    _write(pd_val, f"{ps} ~ {pe} (모니터링 기간 포함)",
           font=FONT_BODY, size=PT_BODY, align=WD_ALIGN_PARAGRAPH.LEFT)
    _set_cell_borders(pd_val,
                      right=SZ_THICK,
                      top=SZ_THIN, left=SZ_THIN, bottom=SZ_THIN)
    _set_row_height(tbl.rows[r_content + 1], 14.66)

    # ── 본문 행 (마지막): 본문 텍스트 + 날짜 + 중첩 서명 표 ──
    r_body = ROWS - 1
    body_cell = tbl.cell(r_body, 0).merge(tbl.cell(r_body, 3))
    _set_cell_borders(body_cell,
                      left=SZ_THICK, right=SZ_THICK, bottom=SZ_THICK,
                      top=SZ_THIN)
    _valign(body_cell, "top")

    body_text = (
        "   협약당사자는 붙임 문서에 의하여 협약을 체결하고, 신의에 따라 성실히 "
        "협약상의 의무를 이행할 것을 확약하여 그 증거로써 본 협약서를 작성하며, "
        "기명날인한 후 각각 1부씩 보관한다.\n(붙임) 협약서 1부.  끝."
    )
    body_cell.text = ""
    para_body = body_cell.paragraphs[0]
    para_body.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _no_spacing(para_body)
    _no_auto_space(para_body)
    run_body = para_body.add_run(body_text)
    _font(run_body, name=FONT_BODY, size=PT_BODY)

    para_date = body_cell.add_paragraph()
    para_date.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _no_spacing(para_date)
    run_date = para_date.add_run(sd)
    _font(run_date, name=FONT_BODY, size=PT_SIGN)

    sign_orgs = [cw, cl] if is_2way else [cw, cl, ce]
    body_cell._tc.append(_build_nested_sign_xml(sign_orgs))

    trailing_p = OxmlElement("w:p")
    body_cell._tc.append(trailing_p)

    _set_row_height(tbl.rows[r_body], 82.70, exact=False)


# ══════════════════════════════════════════════
# 페이지 설정
# ══════════════════════════════════════════════

def _configure_page(doc: Document) -> None:
    sec = doc.sections[0]
    sec.page_width    = Mm(210)
    sec.page_height   = Mm(297)
    sec.left_margin   = L_MARGIN
    sec.right_margin  = R_MARGIN
    sec.top_margin    = T_MARGIN
    sec.bottom_margin = B_MARGIN
    normal = doc.styles["Normal"]
    normal.font.name = FONT_BODY
    normal.font.size = Pt(PT_BODY)


# ══════════════════════════════════════════════
# 공개 API
# ══════════════════════════════════════════════

def make_agreement_2way(data: dict) -> Path:
    """표준 업무협약서 2자용 DOCX 생성."""
    doc = Document()
    _configure_page(doc)
    _build_heading_table(doc, heading_num=3)
    _build_cover_table(doc, data, party_type='2way')
    filename = f"공통3_표준협약서_2자용_{date.today().strftime('%Y%m%d')}.docx"
    return save_docx(doc, filename)


def make_agreement_3way(data: dict) -> Path:
    """표준 업무협약서 3자용 DOCX 생성."""
    doc = Document()
    _configure_page(doc)
    _build_heading_table(doc, heading_num=3)
    _build_cover_table(doc, data, party_type='3way')
    filename = f"공통3_표준협약서_3자용_{date.today().strftime('%Y%m%d')}.docx"
    return save_docx(doc, filename)


def make_agreement(party_type: str = '2way', data: dict | None = None) -> Path:
    """
    표준 업무협약서 생성 통합 함수.
    party_type: '2way' (2자용) 또는 '3way' (3자용)
    """
    if data is None:
        data = {}
    if party_type == '2way':
        return make_agreement_2way(data)
    elif party_type == '3way':
        return make_agreement_3way(data)
    else:
        raise ValueError(f"party_type must be '2way' or '3way', got {party_type!r}")


# ══════════════════════════════════════════════
# 테스트
# ══════════════════════════════════════════════

if __name__ == "__main__":
    SAMPLE = {
        "project_name":    "자활기업 창업 지원",
        "center_wide":     "제주특별자치도광역자활센터",
        "center_local":    "제주시지역자활센터",
        "enterprise_name": "제주자활기업 1호",
        "rep_wide":        "홍길동",
        "rep_local":       "김철수",
        "rep_enterprise":  "이영희",
        "reg_wide":        "123-45-67890",
        "reg_local":       "234-56-78901",
        "reg_enterprise":  "345-67-89012",
        "addr_wide":       "제주시 연동 000-0",
        "addr_local":      "제주시 이도동 000-0",
        "addr_enterprise": "제주시 노형동 000-0",
        "support_content": "사업비 지원",
        "period_start":    "2025.01.01.",
        "period_end":      "2025.12.31.",
        "contract_no":     "제2025-001호",
        "sign_date":       "2025.  01.  01.",
    }

    import subprocess

    p2 = make_agreement("2way", SAMPLE)
    print(f"2자용: {p2}")

    p3 = make_agreement("3way", SAMPLE)
    print(f"3자용: {p3}")

    import sys as _sys
    if _sys.platform == "win32":
        subprocess.Popen(["start", "", str(p2)], shell=True)
        import time; time.sleep(1)
        subprocess.Popen(["start", "", str(p3)], shell=True)
