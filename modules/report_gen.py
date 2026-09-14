"""
report_gen.py
Claude.ai 붙여넣기 방식 보고서 생성 모듈

워크플로:
  1. generate_prompt(data)  → 프롬프트 텍스트 출력
  2. 담당자가 해당 텍스트를 Claude.ai에 붙여넣기
  3. Claude.ai 응답을 복사
  4. make_report_docx(응답_텍스트, data) → DOCX + PDF 저장
"""

import textwrap
from datetime import datetime
from pathlib import Path

from modules.doc_builder import (
    create_doc,
    add_title,
    add_paragraph,
    add_table,
    save_docx,
    docx_to_pdf,
    set_font,
)


def _today_str() -> str:
    return datetime.now().strftime("%Y%m%d")


def _fmt_int(value) -> str:
    """천 단위 쉼표 포함 정수 문자열 반환. 변환 불가 시 원본 반환."""
    try:
        return f"{int(str(value).replace(',', '')):,}"
    except (ValueError, TypeError):
        return str(value)


def _calc_rate(numerator, denominator) -> float:
    """집행률(%) 계산. 분모가 0이면 0.0 반환."""
    try:
        n = int(str(numerator).replace(",", ""))
        d = int(str(denominator).replace(",", ""))
        return round(n / d * 100, 1) if d > 0 else 0.0
    except (ValueError, TypeError, ZeroDivisionError):
        return 0.0


def _add_divider(doc) -> None:
    """수평 구분선을 XML 단락 테두리로 추가한다."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    para = doc.add_paragraph()
    pPr = para._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")        # 0.5pt
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    pBdr.append(bottom)
    pPr.append(pBdr)


# ──────────────────────────────────────────────
# 함수 1: 프롬프트 생성
# ──────────────────────────────────────────────

def generate_prompt(data: dict) -> str:
    """
    Claude.ai에 붙여넣을 보고서 작성 프롬프트를 생성한다.

    Args:
        data: 보고서 작성에 필요한 dict
              필수 키: 사업명, 사업기간, 지원기업수, 완료기업수,
                       지원금총액, 집행금액, 주요성과, 애로사항, 향후계획
    Returns:
        Claude.ai에 붙여넣을 프롬프트 문자열
    """
    집행률 = _calc_rate(data["집행금액"], data["지원금총액"])

    prompt = textwrap.dedent(f"""\
        [보건복지부 제출용 공문서 보고서 작성 요청]

        다음 실적 데이터를 보건복지부 제출용 공문서 형식의 보고서 본문으로 작성해주세요.

        작성 규칙:
        - 문체: 3인칭 기술체, 객관적이고 간결하게
        - 수치: 입력값 그대로 유지
        - 구성: 아래 4개 섹션 순서로 반드시 작성
          1. 사업 개요 (2~3문장)
          2. 추진 실적 (수치 포함 서술)
          3. 예산 집행 현황 (집행률 % 포함)
          4. 성과 및 향후계획

        [실적 데이터]
        사업명: {data['사업명']}
        사업기간: {data['사업기간']}
        지원기업수: {data['지원기업수']}개소
        완료기업수: {data['완료기업수']}개소
        지원금총액: {_fmt_int(data['지원금총액'])}원
        집행금액: {_fmt_int(data['집행금액'])}원
        집행률: {집행률}%
        주요성과: {data['주요성과']}
        애로사항: {data['애로사항']}
        향후계획: {data['향후계획']}

        위 데이터를 바탕으로 보고서 본문만 작성해주세요.
        (제목, 인사말, 부가 설명 없이 본문 텍스트만)
    """)
    return prompt


# ──────────────────────────────────────────────
# 함수 2: Claude.ai 응답 → DOCX
# ──────────────────────────────────────────────

def make_report_docx(prompt_result: str, data: dict) -> Path:
    """
    Claude.ai에서 복사한 응답 텍스트를 받아 DOCX 보고서를 완성한다.

    Args:
        prompt_result: Claude.ai 응답 텍스트 (보고서 본문)
        data:          generate_prompt()에 전달했던 원본 입력 dict
                       필수 키: 사업명, 사업기간, 지원기업수, 완료기업수,
                                지원금총액, 집행금액
    Returns:
        저장된 .docx 파일의 Path 객체
    """
    doc = create_doc()
    집행률 = _calc_rate(data["집행금액"], data["지원금총액"])

    # 1. 제목
    add_title(doc, f"2025년 {data['사업명']} 결과보고서")
    doc.add_paragraph()

    # 2. 기본정보 표
    add_paragraph(doc, "■ 기본 정보", size=10, bold=True)
    info_data = [
        ["항목", "내용"],
        ["사업명",    data["사업명"]],
        ["사업기간",  data["사업기간"]],
        ["지원기업수", f"{data['지원기업수']}개소"],
        ["완료기업수", f"{data['완료기업수']}개소"],
        ["지원금총액", f"{_fmt_int(data['지원금총액'])}원"],
        ["집행금액",  f"{_fmt_int(data['집행금액'])}원"],
        ["집행률",    f"{집행률}%"],
    ]
    add_table(doc, info_data, col_widths=[4.0, 12.5], header=True)
    doc.add_paragraph()

    # 3. 구분선
    _add_divider(doc)

    # 4. 보고서 본문 — 줄바꿈 기준으로 문단 분리
    add_paragraph(doc, "■ 보고서 본문", size=10, bold=True)
    doc.add_paragraph()

    for line in prompt_result.splitlines():
        stripped = line.rstrip()
        if stripped:
            add_paragraph(doc, stripped, size=10)
        else:
            doc.add_paragraph()  # 빈 줄은 빈 문단으로

    doc.add_paragraph()

    # 5. 저장
    filename = f"보고서_{data['사업명']}_{_today_str()}.docx"
    path = save_docx(doc, filename)
    docx_to_pdf(path)
    return path
