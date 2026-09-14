"""
common_forms.py — HWP 원본 서식 기준 재작성
공통3: 표준협약서 / 공통4: 개인정보동의서 / 공통5: 청렴이행각서
"""

from datetime import datetime
from pathlib import Path
from typing import List

import pandas as pd
from docx import Document
from docx.shared import Pt, Cm, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from modules.doc_builder import create_doc, save_docx, docx_to_pdf

FONT_NAME = "맑은 고딕"
DEFAULT_ORG = "제주특별자치도광역자활센터"


# ──────────────────────────────────────────────
# 내부 헬퍼
# ──────────────────────────────────────────────

def _today() -> str:
    return datetime.now().strftime("%Y%m%d")


def _apply_font(run, name=FONT_NAME, size=10, bold=False):
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(attr), name)
    run.font.size = Pt(size)
    run.font.bold = bold


def _para(doc, text="", size=10, bold=False, align="left", spacing=True, indent_cm=0.0):
    """문단 추가 (160% 줄간격 기본)."""
    ALIGN = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }
    p = doc.add_paragraph()
    p.alignment = ALIGN.get(align, WD_ALIGN_PARAGRAPH.LEFT)
    if spacing:
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        p.paragraph_format.line_spacing = 1.6
    if indent_cm:
        p.paragraph_format.left_indent = Cm(indent_cm)
    if text:
        run = p.add_run(text)
        _apply_font(run, size=size, bold=bold)
    return p


def _set_table_borders(table, outer_pt=1.0, inner_pt=0.5):
    """테이블 외곽(outer_pt) / 내부(inner_pt) 테두리 설정."""
    tbl = table._tbl
    tblPr = tbl.find(qn("w:tblPr"))
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
    tblBorders = OxmlElement("w:tblBorders")
    for side, pt in [
        ("top", outer_pt), ("left", outer_pt),
        ("bottom", outer_pt), ("right", outer_pt),
        ("insideH", inner_pt), ("insideV", inner_pt),
    ]:
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(int(pt * 8)))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        tblBorders.append(el)
    tblPr.append(tblBorders)


def _cell_text(cell, text, size=10, bold=False, align="left", bg=None):
    """셀에 텍스트와 스타일을 설정한다."""
    cell.text = ""
    ALIGN = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }
    p = cell.paragraphs[0]
    p.alignment = ALIGN.get(align, WD_ALIGN_PARAGRAPH.LEFT)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(str(text))
    _apply_font(run, size=size, bold=bold)
    if bg:
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), bg)
        tcPr.append(shd)


def _merge_row_label(table, row_idx, label, size=10, bold=True):
    """행 첫 번째 셀을 레이블로 설정 (배경색 D9D9D9)."""
    cell = table.rows[row_idx].cells[0]
    _cell_text(cell, label, size=size, bold=bold, align="center", bg="D9D9D9")


def _add_page_break(doc):
    from docx.oxml import OxmlElement
    p = doc.add_paragraph()
    run = p.add_run()
    run._r.append(OxmlElement("w:br"))
    run._r[-1].set(qn("w:type"), "page")


# ──────────────────────────────────────────────
# 공통3: 표준협약서
# ──────────────────────────────────────────────

_ARTICLES_2WAY = [
    ("제1조(목적)",
     "본 협약은 「국민기초생활보장법」 제18조에 의한 자활기업의 창업 및 경영 안정 지원을 위하여 "
     "광역자활센터(이하 '광역센터'라 한다)와 지역자활센터(이하 '지역센터'라 한다) 간의 "
     "상호 역할과 책임을 규정함을 목적으로 한다."),
    ("제2조(협약의 효력)",
     "이 협약은 협약 체결일로부터 효력이 발생하며, 협약 당사자 쌍방이 서명 또는 날인한 날부터 유효하다."),
    ("제3조(사업내용)",
     "지역센터는 광역센터가 지원하는 자활기업 창업자금을 활용하여 "
     "신청서상의 사업 목적에 맞게 성실히 사업을 수행하여야 한다."),
    ("제4조(지원내용)",
     "광역센터는 지역센터에 아래와 같이 자활기업 창업자금을 지원한다."),
    ("제5조(광역센터 의무)",
     "① 광역센터는 지역센터(자활기업)의 사업 전반에 대하여 지도·감독한다.\n"
     "② 광역센터는 협약에 따른 사업비를 지급하고 집행 상황을 관리한다.\n"
     "③ 광역센터는 사업 운영을 위한 컨설팅 및 기술적 지원을 제공한다.\n"
     "④ 광역센터는 사업기간 중 모니터링을 실시하고 그 결과를 관리한다."),
    ("제6조(지역센터 의무)",
     "① 지역센터는 사업계획서에 따라 성실하게 사업을 수행하여야 한다.\n"
     "② 지역센터는 지원금을 사업 목적 외에 사용하여서는 아니 된다.\n"
     "③ 지역센터는 사업 관련 서류를 보관하고 요청 시 제출하여야 한다.\n"
     "④ 지역센터는 광역센터의 지도·감독에 적극 협력하여야 한다."),
    ("제7조(사업변경)",
     "지역센터는 사업내용 또는 예산을 변경하고자 할 때에는 사전에 광역센터의 서면 승인을 받아야 한다."),
    ("제8조(사업포기)",
     "지역센터가 사업을 포기하고자 할 때에는 즉시 광역센터에 통보하고, "
     "이미 지급받은 지원금 전액을 반환하여야 한다."),
    ("제9조(협약해지)",
     "광역센터는 협약 당사자가 본 협약의 내용을 위반하거나 "
     "사업 목적 달성이 불가능하다고 판단되는 경우 본 협약을 해지할 수 있다."),
    ("제10조(협약효력)",
     "본 협약은 협약기간 종료 후에도 정산, 감사 등을 위하여 5년간 관련 의무사항의 효력이 유지된다."),
    ("제11조(중앙자활자금 정산)",
     "사업 종료 후 잔액이 발생하거나 사업 수행 과정에서 수익금이 발생한 경우, "
     "지역센터는 해당 금액을 광역센터에 반환하여야 한다."),
]

_ARTICLES_3WAY = _ARTICLES_2WAY[:5] + [
    ("제6조(지역센터 및 자활기업 의무)",
     "① 지역센터 및 자활기업은 사업계획서에 따라 성실하게 사업을 수행하여야 한다.\n"
     "② 지역센터 및 자활기업은 지원금을 사업 목적 외에 사용하여서는 아니 된다.\n"
     "③ 지역센터 및 자활기업은 사업 관련 서류를 보관하고 요청 시 제출하여야 한다.\n"
     "④ 지역센터 및 자활기업은 광역센터의 지도·감독에 적극 협력하여야 한다."),
] + _ARTICLES_2WAY[6:]


def make_agreement(party_type: str = "2way", data: dict = None) -> Path:
    """
    표준협약서(공통3)를 생성한다.

    Args:
        party_type: '2way'(광역+지역) 또는 '3way'(광역+지역+자활기업)
        data: 협약서 데이터 dict
    """
    if data is None:
        data = {}
    doc = create_doc()

    # ── 표지 테이블 ──
    _para(doc, "", spacing=False)  # 상단 여백

    # 협약번호 행 (1열 병합)
    cover = doc.add_table(rows=1, cols=2)
    cover.style = "Table Grid"
    _set_table_borders(cover, outer_pt=1.0, inner_pt=0.5)
    cover.rows[0].cells[0].merge(cover.rows[0].cells[1])
    _cell_text(cover.rows[0].cells[0],
               f"협약번호: {data.get('contract_no', '제2025-   호')}",
               size=10, bold=False, align="right")

    # 협약당사자 섹션 레이블 추가
    row = cover.add_row()
    row.cells[0].merge(row.cells[1])
    _cell_text(row.cells[0], "협 약 당 사 자", size=11, bold=True, align="center", bg="D9D9D9")

    # 광역 정보
    for label, key in [("광역자활센터명", "center_wide"),
                        ("대표자", "rep_wide"),
                        ("사업자등록번호", "reg_wide")]:
        r = cover.add_row()
        _cell_text(r.cells[0], label, size=10, bold=True, align="center", bg="EFEFEF")
        _cell_text(r.cells[1], data.get(key, ""), size=10)

    # 지역 정보
    for label, key in [("지역자활센터명", "center_local"),
                        ("대표자", "rep_local"),
                        ("사업자등록번호", "reg_local")]:
        r = cover.add_row()
        _cell_text(r.cells[0], label, size=10, bold=True, align="center", bg="EFEFEF")
        _cell_text(r.cells[1], data.get(key, ""), size=10)

    # 3way: 자활기업 추가
    if party_type == "3way":
        for label, key in [("자활기업명", "enterprise_name"),
                            ("대표자", "rep_enterprise"),
                            ("사업자등록번호", "reg_enterprise")]:
            r = cover.add_row()
            _cell_text(r.cells[0], label, size=10, bold=True, align="center", bg="EFEFEF")
            _cell_text(r.cells[1], data.get(key, ""), size=10)

    # 협약내용 섹션
    r = cover.add_row()
    r.cells[0].merge(r.cells[1])
    _cell_text(r.cells[0], "협 약 내 용", size=11, bold=True, align="center", bg="D9D9D9")

    for label, val in [
        ("지원내용", data.get("support_content", "")),
        ("협약기간", f"{data.get('period_start','')} ~ {data.get('period_end','')} (모니터링 기간 포함)"),
    ]:
        r = cover.add_row()
        _cell_text(r.cells[0], label, size=10, bold=True, align="center", bg="EFEFEF")
        _cell_text(r.cells[1], val, size=10)

    # ── 표지 서명란 ──
    _para(doc)
    _para(doc, "위와 같이 협약을 체결합니다.", size=10, align="center")
    _para(doc)
    _para(doc, "2025년    월    일", size=10, align="center")
    _para(doc)

    wide_label = f"{data.get('center_wide', '○○광역자활센터')}"
    local_label = f"{data.get('center_local', '○○지역자활센터')}"

    if party_type == "2way":
        sign_tbl = doc.add_table(rows=3, cols=2)
        sign_tbl.style = "Table Grid"
        _set_table_borders(sign_tbl)
        _cell_text(sign_tbl.rows[0].cells[0], wide_label, bold=True, align="center")
        _cell_text(sign_tbl.rows[0].cells[1], local_label, bold=True, align="center")
        _cell_text(sign_tbl.rows[1].cells[0], f"대표자: {data.get('rep_wide','')}", align="center")
        _cell_text(sign_tbl.rows[1].cells[1], f"대표자: {data.get('rep_local','')}", align="center")
        _cell_text(sign_tbl.rows[2].cells[0], "(직 인)", size=10, align="center")
        _cell_text(sign_tbl.rows[2].cells[1], "(직 인)", size=10, align="center")
    else:
        ent_label = f"{data.get('enterprise_name', '○○자활기업')}"
        sign_tbl = doc.add_table(rows=3, cols=3)
        sign_tbl.style = "Table Grid"
        _set_table_borders(sign_tbl)
        for i, (lbl, rep_key) in enumerate([
            (wide_label, "rep_wide"),
            (local_label, "rep_local"),
            (ent_label, "rep_enterprise"),
        ]):
            _cell_text(sign_tbl.rows[0].cells[i], lbl, bold=True, align="center")
            _cell_text(sign_tbl.rows[1].cells[i], f"대표자: {data.get(rep_key,'')}", align="center")
            _cell_text(sign_tbl.rows[2].cells[i], "(직 인)", size=10, align="center")

    # ── 본문 페이지 시작 ──
    _add_page_break(doc)
    _para(doc, "협  약  서  본  문", size=14, bold=True, align="center", spacing=False)
    _para(doc)

    articles = _ARTICLES_3WAY if party_type == "3way" else _ARTICLES_2WAY

    for title, body in articles:
        _para(doc, title, size=10, bold=True)

        # 제4조: 예산 테이블 삽입
        if "제4조" in title:
            _para(doc, body, size=10)
            budget_tbl = doc.add_table(rows=2, cols=5)
            budget_tbl.style = "Table Grid"
            _set_table_borders(budget_tbl)
            headers = ["구분", "중앙자활자금", "자부담", "자활기금", "지역자활사업지원비"]
            values = [
                "금액(원)",
                data.get("amount_central", ""),
                data.get("amount_self", ""),
                data.get("amount_fund", ""),
                data.get("amount_local_support", ""),
            ]
            for i, (h, v) in enumerate(zip(headers, values)):
                _cell_text(budget_tbl.rows[0].cells[i], h, bold=True, align="center", bg="D9D9D9")
                _cell_text(budget_tbl.rows[1].cells[i], str(v), align="center")
            col_w = [Cm(2.5), Cm(3.5), Cm(2.5), Cm(2.5), Cm(4.0)]
            for i, w in enumerate(col_w):
                for row in budget_tbl.rows:
                    row.cells[i].width = w
        else:
            # 줄바꿈 포함 본문
            for line in body.split("\n"):
                indent = 0.5 if line.startswith("①") or line.startswith("②") or \
                               line.startswith("③") or line.startswith("④") else 0.0
                _para(doc, line, size=10, indent_cm=indent)

        _para(doc, "", spacing=False)

    # ── 본문 서명란 ──
    _para(doc)
    _para(doc, "2025년    월    일", size=10, align="center")
    _para(doc)

    if party_type == "2way":
        _para(doc, f"{wide_label}  대표자: {data.get('rep_wide','')} (서명 또는 인)", size=10, align="right")
        _para(doc, f"{local_label}  대표자: {data.get('rep_local','')} (서명 또는 인)", size=10, align="right")
    else:
        _para(doc, f"{wide_label}  대표자: {data.get('rep_wide','')} (서명 또는 인)", size=10, align="right")
        _para(doc, f"{local_label}  대표자: {data.get('rep_local','')} (서명 또는 인)", size=10, align="right")
        _para(doc, f"{ent_label}  대표자: {data.get('rep_enterprise','')} (서명 또는 인)", size=10, align="right")

    filename = f"공통3_표준협약서_{party_type}_{data.get('center_local','기관')}_{_today()}.docx"
    path = save_docx(doc, filename)
    docx_to_pdf(path)
    return path


# ──────────────────────────────────────────────
# 공통4: 개인정보 수집·이용 동의서
# ──────────────────────────────────────────────

def make_privacy_consent(data: dict = None) -> Path:
    """개인정보 수집·이용 동의서(공통4)를 생성한다."""
    if data is None:
        data = {}
    doc = create_doc()

    # 1. 제목
    _para(doc, "개인정보 수집 및 이용 동의서", size=16, bold=True, align="center", spacing=False)
    _para(doc, f"({data.get('project_name', '중앙자활자금 지원사업')})", size=11, align="center", spacing=False)
    _para(doc)

    # 2. 안내문
    _para(doc,
          "아래와 같이 개인정보를 수집 및 이용하고자 하오니,\n"
          "자세히 읽어보신 후 동의여부를 결정하여 주시기 바랍니다.",
          size=10)
    _para(doc)

    # 3. 필수항목 섹션
    _para(doc, "○ 필수항목", size=10, bold=True)

    items = [
        ("수집항목\n(광역자활센터)",
         "성명, 연락처, 생년월일, 자격 및 경력사항(자격증 및 경력증명서 사본),\n"
         "고용보험 가입이력, 가족관계(가족관계증명서)"),
        ("수집항목\n(한국자활복지개발원)",
         "성명, 연락처, 생년월일, 자격 및 경력사항"),
        ("수집목적",
         "중앙자활자금 지원사업 참여를 위한 채용 승인 검토 및 사업 관리"),
        ("보유기간", "수집일로부터 5년"),
        ("거부 시\n불이익",
         "위와 같이 개인정보를 수집 및 이용하는데 동의를 거부할 권리가 있습니다.\n"
         "동의를 거부할 경우 중앙자활자금 지원사업 선정에 제외될 수 있습니다."),
    ]

    tbl = doc.add_table(rows=len(items), cols=2)
    tbl.style = "Table Grid"
    _set_table_borders(tbl)

    col_widths = [Cm(4.0), Cm(12.5)]
    for row_idx, (label, content) in enumerate(items):
        row = tbl.rows[row_idx]
        row.cells[0].width = col_widths[0]
        row.cells[1].width = col_widths[1]
        _cell_text(row.cells[0], label, size=10, bold=True, align="center", bg="D9D9D9")
        _cell_text(row.cells[1], content, size=10)

    _para(doc)

    # 4. 동의 여부
    _para(doc,
          "위와 같이 개인정보를 수집 및 이용하는 것에 동의하십니까?",
          size=10, bold=True)
    _para(doc, "□ 동의함          □ 동의하지 않음", size=10, align="center")
    _para(doc)

    # 5. 서명란
    sign_year = data.get("sign_date", "2025")
    _para(doc, f"{sign_year}년     월     일", size=10, align="center")
    _para(doc)
    _para(doc, "이    름 :                              (날인)", size=10, align="center")
    _para(doc)

    # 6. 수신
    center_wide = data.get("center_wide", DEFAULT_ORG)
    _para(doc, f"한국자활복지개발원장 · {center_wide}장 귀하", size=10, align="right")

    filename = f"공통4_개인정보동의서_{data.get('signee_name','동의자')}_{_today()}.docx"
    path = save_docx(doc, filename)
    docx_to_pdf(path)
    return path


# ──────────────────────────────────────────────
# 공통5: 청렴이행각서
# ──────────────────────────────────────────────

def make_integrity_pledge(data: dict = None) -> Path:
    """청렴이행각서(공통5)를 생성한다."""
    if data is None:
        data = {}
    doc = create_doc()
    center_wide = data.get("center_wide", DEFAULT_ORG)

    # 1. 제목
    _para(doc, "청  렴  이  행  각  서", size=16, bold=True, align="center", spacing=False)
    _para(doc)

    # 2. 서문
    _para(doc,
          f"우리 기관은 부패없는 투명하고 공정한 행정이 사회발전과 국가경쟁력에 중요한 관건이 됨을 "
          f"깊이 인식하고, 청렴이행각서 작성 취지에 적극 호응하며 "
          f"한국자활복지개발원({center_wide})에서 실시하는 지원사업을 참여함에 있어 "
          f"우리 기관의 직원과 대리인은",
          size=10)
    _para(doc)

    # 3. 조항 3개
    clauses = [
        "1. 한국자활복지개발원({center_wide}) 및 관련기관 임직원에게 "
        "금품·향응 등을 제공하지 않겠으며, 위반 시 향후 3년간 관련 지원사업 참여를 "
        "제한 받는 것에 동의합니다.",
        "2. 상기와 같은 행위로 선정이 취소되거나 협약이 해제되더라도 이에 대한 "
        "민·형사상 일체의 이의를 제기하지 않겠습니다.",
        "3. 지원된 예산은 목적 외에 사용하지 않겠으며, 위반 사항이 발생할 경우 "
        "반납조치를 감수하겠습니다.",
    ]
    for clause in clauses:
        _para(doc, clause.format(center_wide=center_wide), size=10, indent_cm=0.5)

    _para(doc)

    # 4. 서약 문구
    _para(doc,
          "위 청렴이행각서는 상호신뢰를 바탕으로 한 약속으로서 반드시 지킬 것이며, "
          "이를 위반할 경우 어떠한 제재도 감수할 것을 서약합니다.",
          size=10)
    _para(doc)

    # 5. 서명란
    sign_year = data.get("sign_date", "2025")
    _para(doc, f"{sign_year}년     월     일", size=10, align="center")
    _para(doc)
    org_name = data.get("org_name", "")
    rep_name = data.get("rep_name", "")
    _para(doc, f"서약기관명: {org_name}   대표 {rep_name} (서명 또는 인)", size=10, align="right")
    _para(doc)

    # 6. 수신
    _para(doc, f"{center_wide}장 귀중", size=10, align="right")

    filename = f"공통5_청렴이행각서_{data.get('org_name','기관')}_{_today()}.docx"
    path = save_docx(doc, filename)
    docx_to_pdf(path)
    return path


# ──────────────────────────────────────────────
# 일괄 생성
# ──────────────────────────────────────────────

def make_all_common(csv_path: str, form_type: str) -> List[Path]:
    """
    CSV 일괄 생성.
    form_type: '협약서' / '동의서' / '각서'
    """
    csv_path = Path(csv_path)
    if not csv_path.exists():
        raise FileNotFoundError(f"CSV 파일 없음: {csv_path}")

    handlers = {
        "협약서": lambda d: make_agreement(d.get("party_type", "2way"), d),
        "동의서": make_privacy_consent,
        "각서":   make_integrity_pledge,
    }
    if form_type not in handlers:
        raise ValueError(f"지원하지 않는 서식: '{form_type}'. 가능: {list(handlers)}")

    df = pd.read_csv(csv_path, encoding="utf-8-sig", dtype=str).fillna("")
    results: List[Path] = []
    for idx, row in df.iterrows():
        try:
            path = handlers[form_type](row.to_dict())
            results.append(path)
            print(f"[완료] {path.name}")
        except Exception as e:
            print(f"[경고] {idx+1}행 실패: {e}")
    return results
